# -*- coding: utf-8 -*-
"""ประกอบเด็ค .pptx จากสไลด์ไบล์พรินต์ (Trainer_Docs/Slide_Blueprints/*.md)
บน template Trainocate_SQLPerfTuning_Template (8 layouts) — ใส่ข้อความผ่าน placeholder เท่านั้น
ตามมาตรฐานสไลด์ 4.1–4.6 ของแผนแม่บท"""
import re
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

REPO = Path(__file__).resolve().parents[1]
BP = Path(__file__).parent / "Slide_Blueprints"
OUT = Path(__file__).parent / "Decks"
OUT.mkdir(exist_ok=True)
TPL = Path(__file__).parent / "_template_work.pptx"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
DARK = {"01 Cover (dark)", "02 Module Divider (dark)", "08 Closing (dark)"}
THAI = "Leelawadee UI"

def parse(md_text):
    slides, cur, field = [], None, None
    for line in md_text.splitlines():
        m = re.match(r"^### สไลด์ (\d+): (.*)$", line)
        if m:
            cur = {"num": int(m.group(1)), "title": m.group(2).strip(), "bullets": [], "meta": {}}
            slides.append(cur)
            field = None
            continue
        if cur is None:
            continue
        fm = re.match(r"^- \*\*(Layout|Bullets[^*]*|เนื้อหา[^*]*|ภาพ/แผนภาพ|Scripts?|Speaker notes):\*\* ?(.*)$", line)
        if fm:
            field = fm.group(1)
            if field.startswith("Bullets") or field.startswith("เนื้อหา"):
                field = "Bullets"
            cur["meta"][field] = fm.group(2).strip()
            continue
        bm = re.match(r"^\s+(?:[-•]|\d+[.)])\s+(.*)$", line)
        if bm and field == "Bullets":
            cur["bullets"].append(bm.group(1).strip())
    return slides

def layout_name(s):
    lay = s["meta"].get("Layout", "")
    if "01 Cover" in lay: return "01 Cover (dark)"
    if "section-divider" in lay: return "02 Module Divider (dark)"
    if "two-column" in lay: return "04 Two-Column"
    if "table" in lay: return "05 Table"
    if "demo" in lay or "lab" in lay: return "06 Code Lab"
    if "stat" in lay: return "07 Stat Callout"
    if "summary" in lay: return "08 Closing (dark)"
    return "03 Content"

TITLE_STRIP = re.compile(r"^(Cover|Module Divider|Section Divider|Demo|Lab 1|Closing|Quiz)\s*—\s*")

def clean_title(s):
    return TITLE_STRIP.sub("", s["title"]).strip()

def inline(text):
    """'**bold** `code`' -> [(text, bold, mono)]"""
    out = []
    for seg in re.split(r"(\*\*[^*]+\*\*)", text):
        if seg.startswith("**") and seg.endswith("**"):
            out.append((seg[2:-2].replace("`", ""), True, False))
        else:
            for j, sub in enumerate(re.split(r"(`[^`]+`)", seg)):
                if sub.startswith("`") and sub.endswith("`"):
                    out.append((sub[1:-1], False, True))
                elif sub:
                    out.append((sub, False, False))
    return out or [("", False, False)]

def fill_para(p, spec, size, bullet=True, space_after=8, light_bg=True):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "228600" if bullet else "0")
    pPr.set("indent", "-228600" if bullet else "0")
    if bullet:
        etree.SubElement(pPr, f"{{{A}}}buFont").set("typeface", "Arial")
        etree.SubElement(pPr, f"{{{A}}}buChar").set("char", "▪")
    else:
        etree.SubElement(pPr, f"{{{A}}}buNone")
    for text, bold, mono in spec:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = "Consolas" if mono else THAI
        if light_bg:
            r.font.color.rgb = RGBColor(0x24, 0x30, 0x40)
        rPr = r._r.get_or_add_rPr()
        etree.SubElement(rPr, f"{{{A}}}cs").set("typeface", "Consolas" if mono else THAI)
    p.space_after = Pt(space_after)

def body_of(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx != 0:
            return ph
    return None

def title_of(slide):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            return ph
    return None

def add_notes(slide, s):
    notes = s["meta"].get("Speaker notes", "")
    notes = notes.replace("Trainer_Docs/Glossary.md (สร้างในเฟส 0 — วันนี้ยังไม่มีไฟล์ใน repo)", "Glossary.md (repo root)")
    extra = s.get("extra_notes") or []
    if extra:
        notes += "\nจุดเพิ่มเติมที่ย่อออกจากสไลด์: " + " | ".join(extra)
    if "Ref:" not in notes:
        notes += " Ref: " + clean_title(s)
    scripts = s["meta"].get("Scripts", "")
    if scripts:
        notes += "\nScripts: " + scripts
    slide.notes_slide.notes_text_frame.text = notes

def split_cols(bullets):
    L, R = [], []
    for b in bullets:
        if re.match(r"^(ซ้าย|ฝั่งซ้าย|\(ซ้าย\))\s*[-—:]?", b):
            L.append(re.sub(r"^(ซ้าย|ฝั่งซ้าย|\(ซ้าย\))\s*[-—:]?\s*", "", b))
        elif re.match(r"^(ขวา|ฝั่งขวา|\(ขวา\))", b):
            R.append(re.sub(r"^(ขวา|ฝั่งขวา|\(ขวา\))(\s*\(ต่อ\))?\s*[-—:]?\s*:?\s*", "", b))
        elif re.match(r"^(\(แถบล่าง\)|แถบล่าง)", b):
            R.append(re.sub(r"^(\(แถบล่าง\)|แถบล่าง)\s*[-—:]?\s*:?\s*", "", b))
        else:
            (L if len("".join(L)) <= len("".join(R)) else R).append(b)
    return L, R

def find_image(s):
    """หาไฟล์ .png ที่ meta 'ภาพ/แผนภาพ' อ้างถึง — คืน path ถ้ามีจริงใน repo"""
    m = re.search(r"([A-Za-z0-9_-]+\.png)", s["meta"].get("ภาพ/แผนภาพ", ""))
    if not m:
        return None, None
    name = m.group(1)
    hit = list(REPO.glob(f"Module_*/Sections/*/images/{name}"))
    return (hit[0], name) if hit else (None, name)


def build_deck(blueprint_path, out_path, only_nums=None, use_images=True):
    md = blueprint_path.read_text(encoding="utf-8")
    slides = parse(md)
    if only_nums:
        slides = [s for s in slides if s["num"] in only_nums]
    prs = Presentation(TPL)
    lay = {l.name: l for l in prs.slide_layouts}
    for s in slides:
        lname = layout_name(s)
        sl = prs.slides.add_slide(lay[lname])
        light = lname not in DARK and lname != "06 Code Lab"
        tp, bd = title_of(sl), body_of(sl)
        if tp is not None:
            tp.text_frame.paragraphs[0].text = ""
            fill_para(tp.text_frame.paragraphs[0], [(clean_title(s), True, False)], 30, bullet=False,
                      light_bg=light)
        size = 15 if len(s["bullets"]) >= 5 else 16
        # ตัดสินใจเรื่องภาพก่อน (ใช้ทั้ง layout และ anchor)
        img, img_name = find_image(s)
        has_img = use_images and img is not None
        if lname == "02 Module Divider (dark)":
            keep = s["bullets"][:2] if s["num"] > 2 else s["bullets"][:3]
            extra = s["bullets"][len(keep):]
            lines = [" · ".join(keep[:2])]
            if len(keep) > 2:
                lines.append(keep[2])
            bd.text_frame.clear()
            for i, ln in enumerate(lines):
                p = bd.text_frame.paragraphs[0] if i == 0 else bd.text_frame.add_paragraph()
                fill_para(p, inline(ln), 15, bullet=False, space_after=4, light_bg=light)
            if extra:
                s.setdefault("extra_notes", [])
                s["extra_notes"] += extra
        elif lname == "04 Two-Column":
            L, R = split_cols(s["bullets"])
            cols = sorted([ph for ph in sl.placeholders if ph.placeholder_format.idx != 0], key=lambda p: p.left)
            from pptx.enum.text import MSO_ANCHOR
            for phx, items in ((cols[0], L), (cols[1], R)):
                phx.text_frame.clear()
                phx.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE  # สมดุลแนวตั้งสองคอลัมน์
                for i, b in enumerate(items):
                    p = phx.text_frame.paragraphs[0] if i == 0 else phx.text_frame.add_paragraph()
                    fill_para(p, inline(b), 15, space_after=8, light_bg=light)
        else:
            bd.text_frame.clear()
            from pptx.enum.text import MSO_ANCHOR
            if not has_img:
                bd.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE  # กันการ์ดโล่ง: เนื้อหาน้อยให้อยู่กึ่งกลาง
            body_size = 13 if (has_img and len(s["bullets"]) >= 5) else size
            for i, b in enumerate(s["bullets"]):
                p = bd.text_frame.paragraphs[0] if i == 0 else bd.text_frame.add_paragraph()
                fill_para(p, inline(b), body_size, space_after=8, light_bg=light)
            if lname == "06 Code Lab" and s["meta"].get("Scripts"):
                p = bd.text_frame.add_paragraph()
                fill_para(p, [("Scripts: " + s["meta"]["Scripts"], False, True)], 11.5, bullet=False,
                          space_after=0, light_bg=light)
        # วางภาพ (ถ้ามีไฟล์จริง) — แคบ body ฝั่งซ้ายแล้ววางภาพฝั่งขวา; ยังไม่มีไฟล์: แจ้งใน notes
        if has_img:
            from PIL import Image
            with Image.open(img) as im:
                ratio = im.width / im.height
            bd.width = Inches(5.90)
            iw = 5.95
            ih = iw / ratio
            if ih > 4.70:
                ih = 4.70
                iw = ih * ratio
            sl.shapes.add_picture(str(img), Inches(6.88), Inches(1.95 + (4.90 - ih) / 2),
                                  Inches(iw), Inches(ih))
        elif use_images and img_name is not None:
            s.setdefault("extra_notes", []).append(
                f"ภาพ {img_name} ยังไม่ถูกวาด (ทีมภาพจะส่งมอบภายหลัง) — สไลด์นี้ยังไม่มีแผนภาพประกอบ")
        add_notes(sl, s)
    prs.save(out_path)
    return len(slides)

def main():
    n = build_deck(BP / "Module_00_Opening_Slides.md", OUT / "00_Opening.pptx")
    print(f"00_Opening.pptx: {n} slides")
    for md in sorted(BP.glob("Module_*_Slides.md")):
        num = md.name.split("_")[1]
        if num in ("00", "01"):
            continue  # 00 สร้างเป็น 00_Opening.pptx แล้ว / 01 ผ่าน visual QA แล้ว
        n = build_deck(md, OUT / f"{num}.pptx")
        print(f"{num}.pptx: {n} slides")

if __name__ == "__main__":
    main()
